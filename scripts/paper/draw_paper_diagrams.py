"""Publication-size Figures 1--3; native vectors and recorded experimental images.

Coordinates use inches from the upper left. The optional PowerPoint exporter
uses the same geometry, and is never imported during manuscript generation.
"""
from pathlib import Path
import json

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, Rectangle, Polygon, PathPatch
from matplotlib.path import Path as MplPath
import numpy as np
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
FIG = ROOT / 'docs/Thesis/ral/fig'
DATA = ROOT / 'logs/paper_ral_20261007'
ASSET = FIG / 'overview_assets'
INK, MUTED, RULE = '#20282C', '#59636A', '#C5CBCF'
TEAL, BLUE = '#167568', '#376687'
SCENES = [f'office{i}' for i in range(5)] + [f'room{i}' for i in range(3)]
TEASER_ROI = (76, 84, 300, 112)


class Drawing:
    def __init__(self, width, height):
        self.width, self.height = width, height
        self.items = []

    def text(self, x, y, value, size: float=8, color=INK, bold=False, align='left',
             ppt=None, runs=None):
        self.items.append(dict(kind='text', x=x, y=y, value=value, size=size,
                               color=color, bold=bold, align=align, ppt=ppt, runs=runs))

    def rect(self, x, y, w, h, color=INK, fill='white', lw=.7, zorder=2):
        self.items.append(dict(kind='rect', x=x, y=y, w=w, h=h,
                               color=color, fill=fill, lw=lw, zorder=zorder))

    def ellipse(self, x, y, w, h, angle=0, color=TEAL, fill='white', lw=.7):
        self.items.append(dict(kind='ellipse', x=x, y=y, w=w, h=h, angle=angle,
                               color=color, fill=fill, lw=lw))

    def line(self, points, color=INK, lw=.8, arrow=False):
        self.items.append(dict(kind='line', points=points, color=color, lw=lw,
                               arrow=arrow))

    def image(self, path, x, y, w, h=None, crop=None):
        with Image.open(path) as img:
            iw, ih = (crop[2], crop[3]) if crop else img.size
        h = h if h is not None else w * ih / iw
        self.items.append(dict(kind='image', path=Path(path), x=x, y=y, w=w,
                               h=h, crop=crop))
        return h

    def heading(self, x, y, label, reference=None):
        self.text(x, y, label, size=9, bold=True)
        if reference:
            self.text(x, y+.21, reference, size=7.5, color=MUTED)

    def render(self, name):
        with plt.rc_context({'font.family': 'DejaVu Sans',
                             'mathtext.fontset': 'dejavusans', 'pdf.fonttype': 42,
                             'ps.fonttype': 42, 'svg.fonttype': 'none'}):
            fig = plt.figure(figsize=(self.width, self.height), facecolor='white')
            ax = fig.add_axes([0, 0, 1, 1])
            ax.set(xlim=(0, self.width), ylim=(self.height, 0))
            ax.axis('off')
            labels = []
            for o in self.items:
                k = o['kind']
                if k == 'text':
                    labels.append(ax.text(o['x'], o['y'], o['value'],
                        fontsize=o['size'], color=o['color'],
                        fontweight='bold' if o['bold'] else 'normal',
                        ha=o['align'], va='center', linespacing=1.22, zorder=4))
                elif k == 'rect':
                    ax.add_patch(Rectangle((o['x'], o['y']), o['w'], o['h'],
                        facecolor=o['fill'], edgecolor=o['color'],
                        linewidth=o['lw'], zorder=o.get('zorder', 2)))
                elif k == 'ellipse':
                    ax.add_patch(Ellipse((o['x'], o['y']), o['w'], o['h'],
                        angle=o['angle'], edgecolor=o['color'],
                        facecolor=o['fill'], linewidth=o['lw'], zorder=3))
                elif k == 'line':
                    pts = np.array(o['points'], dtype=float)
                    ax.add_patch(PathPatch(MplPath(pts,
                        [MplPath.MOVETO]+[MplPath.LINETO]*(len(pts)-1)),
                        fill=False, edgecolor=o['color'], linewidth=o['lw'],
                        capstyle='butt', joinstyle='miter', zorder=3))
                    if o['arrow']:
                        tip = pts[-1]
                        direction = tip - pts[-2]
                        direction /= np.linalg.norm(direction)
                        normal = np.array([-direction[1], direction[0]])
                        base = tip - direction * .056
                        ax.add_patch(Polygon([tip, base+normal*.024, base-normal*.024],
                            facecolor=o['color'], edgecolor=o['color'],
                            linewidth=.3, zorder=3))
                else:
                    with Image.open(o['path']) as img:
                        if o['crop']:
                            cx, cy, cw, ch = o['crop']
                            img = img.crop((cx, cy, cx+cw, cy+ch))
                        ax.imshow(np.array(img), extent=(o['x'], o['x']+o['w'],
                            o['y']+o['h'], o['y']), aspect='auto',
                            interpolation='nearest', zorder=1)
            fig.canvas.draw()
            renderer = fig.canvas.get_renderer()
            for label in labels:
                bb = label.get_window_extent(renderer)
                if not (bb.x0 >= -1 and bb.y0 >= -1 and
                        bb.x1 <= fig.bbox.width+1 and bb.y1 <= fig.bbox.height+1):
                    raise ValueError(f'{name}: text outside page: {label.get_text()}')
            FIG.mkdir(parents=True, exist_ok=True)
            for ext in ('pdf', 'svg', 'png'):
                fig.savefig(FIG / f'{name}.{ext}', dpi=240)
            plt.close(fig)

    def add_slide(self, prs, x=.25, y=.3, width=None):
        from matplotlib.font_manager import FontProperties
        from matplotlib.textpath import TextPath
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Inches, Pt
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        scale = (width or (prs.slide_width/914400-2*x)) / self.width
        def xy(px, py):
            return Inches(x+px*scale), Inches(y+py*scale)
        def color(v):
            return RGBColor.from_string(v.lstrip('#'))
        def no_effects(shape):
            shape._element.spPr.append(OxmlElement('a:effectLst'))
            style = shape._element.find(
                '{http://schemas.openxmlformats.org/presentationml/2006/main}style')
            if style is not None:
                shape._element.remove(style)
        for o in self.items:
            k = o['kind']
            if k == 'text':
                value = o['ppt'] or o['value']
                font = FontProperties(family='DejaVu Sans',
                                      weight='bold' if o['bold'] else 'normal')
                tw = max(.28, max(TextPath((0, 0), s, size=o['size'], prop=font)
                                 .get_extents().width / 72
                                 for s in value.splitlines()) + .05)
                th = len(value.splitlines())*o['size']/72*1.36
                left = o['x'] - (tw/2 if o['align']=='center' else
                                  tw if o['align']=='right' else 0)
                shape = slide.shapes.add_textbox(*xy(left, o['y']-th/2),
                    Inches(tw*scale), Inches(th*scale))
                tf = shape.text_frame
                tf.margin_left = tf.margin_right = 0
                tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                for i, line in enumerate(value.splitlines()):
                    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
                    p.alignment = {'left':PP_ALIGN.LEFT, 'right':PP_ALIGN.RIGHT,
                                   'center':PP_ALIGN.CENTER}[o['align']]
                    p.space_before = p.space_after = Pt(0)
                    for content, baseline in o['runs'] or [(line, 0)]:
                        run = p.add_run()
                        run.text = content
                        run.font.name = 'DejaVu Sans'
                        run.font.size = Pt(o['size']*scale*(.7 if baseline else 1))
                        run.font.bold = o['bold']
                        run.font.color.rgb = color(o['color'])
                        if baseline:
                            run._r.get_or_add_rPr().set('baseline', str(baseline))
            elif k in ('rect', 'ellipse'):
                px, py = o['x'], o['y']
                if k=='ellipse':
                    px -= o['w']/2
                    py -= o['h']/2
                shape = slide.shapes.add_shape(
                    MSO_AUTO_SHAPE_TYPE.RECTANGLE if k=='rect' else
                    MSO_AUTO_SHAPE_TYPE.OVAL, *xy(px, py),
                    Inches(o['w']*scale), Inches(o['h']*scale))
                if o['fill']=='none':
                    shape.fill.background()
                else:
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = color(o['fill']) if o['fill']!='white' else RGBColor(255,255,255)
                shape.line.color.rgb = color(o['color'])
                shape.line.width = Pt(o['lw']*scale)
                if k=='ellipse':
                    shape.rotation = o['angle']
                no_effects(shape)
            elif k=='line':
                points = [xy(px, py) for px, py in o['points']]
                builder = slide.shapes.build_freeform(*points[0])
                builder.add_line_segments(points[1:], close=False)
                shape = builder.convert_to_shape()
                shape.fill.background()
                shape.line.color.rgb = color(o['color'])
                shape.line.width = Pt(o['lw']*scale)
                if o['arrow']:
                    end = OxmlElement('a:tailEnd')
                    end.set('type','triangle'); end.set('w','sm'); end.set('len','sm')
                    shape._element.spPr.get_or_add_ln().append(end)
                no_effects(shape)
            else:
                shape = slide.shapes.add_picture(str(o['path']), *xy(o['x'],o['y']),
                    Inches(o['w']*scale), Inches(o['h']*scale))
                if o['crop']:
                    with Image.open(o['path']) as img:
                        iw, ih = img.size
                    cx, cy, cw, ch = o['crop']
                    shape.crop_left, shape.crop_right = cx/iw, (iw-cx-cw)/iw
                    shape.crop_top, shape.crop_bottom = cy/ih, (ih-cy-ch)/ih
        return slide


def camera(d, x, y, scale=1, color=BLUE):
    d.line([(x,y),(x+.26*scale,y-.14*scale),(x+.26*scale,y+.14*scale),(x,y)],color,.8)
    d.line([(x+.26*scale,y-.14*scale),(x+.40*scale,y-.20*scale),
            (x+.40*scale,y+.20*scale),(x+.26*scale,y+.14*scale)],color,.6)


def gaussians(d, x, y, scale=1, color=TEAL):
    for dx,dy,angle in [(-.15,-.05,-25),(.03,-.13,20),(.19,-.01,-15),(-.04,.08,18)]:
        d.ellipse(x+dx*scale,y+dy*scale,.26*scale,.105*scale,angle,color=color)


def teaser(records):
    d = Drawing(7.15,2.23)
    d.heading(.06,.14,'(a) Shared anchor')
    d.heading(2.27,.14,'(b) Same-camera reconstruction')
    d.heading(5.61,.14,'(c) Replica')
    d.line([(2.13,.05),(2.13,2.15)],RULE,.5)
    d.line([(5.49,.05),(5.49,2.15)],RULE,.5)
    for x,y,muted in [(.30,.80,True),(1.20,.56,False)]:
        c = '#9DA8AE' if muted else INK
        d.line([(x,y),(x+.43,y+.31)],c,.7)
        d.line([(x,y),(x-.10,y+.78)],c,.7)
        d.rect(x-.045,y-.045,.09,.09,c,'white')
        camera(d,x-.10,y+.78,.72,c if muted else BLUE)
        gaussians(d,x+.43,y+.31,.66,c if muted else TEAL)
        d.text(x-.09,y-.10,'k' if muted else 'k′',8,c,align='right')
    d.line([(.345,.788),(1.155,.572)],INK,1,True)
    d.text(.08,.40,'Pose correction',8)
    d.text(1.66,1.16,'local map',7.5,TEAL,align='center')
    d.text(1.35,1.67,'camera',7.5,BLUE,align='center')
    d.text(1.04,2.00,'Map + cameras move together',7.7,bold=True,align='center')
    roi = TEASER_ROI
    for x,m,label in [(2.27,'photo','Photo-SLAM'),(3.93,'ours','Splatt3R-SLAM')]:
        p = DATA/'office0'/f'view0_{m}.png'
        d.text(x+.73,.43,label,8,TEAL if m=='ours' else INK,bold=m=='ours',align='center')
        d.image(p,x,.58,1.46)
        d.rect(x+roi[0]/512*1.46,.58+roi[1]/512*1.46,
               roi[2]/512*1.46,roi[3]/512*1.46,TEAL,'none',.65)
        d.image(p,x,1.50,1.46,crop=roi)
    d.text(3.82,2.14,'Identical crop · representative office0 view',7.3,MUTED,align='center')
    mean = {m:float(np.mean([records[s]['methods'][m]['full_sh']['psnr']
                             for s in SCENES])) for m in ('photo','ours')}
    d.text(5.65,.43,'Mean PSNR (dB) ↑',7.8)
    x0,y0,ch = 5.91,1.69,.99
    for value in (0,10,20):
        y = y0-ch*value/27
        d.line([(x0,y),(7.05,y)],RULE,.45)
        d.text(x0-.07,y,str(value),7.4,MUTED,align='right')
    for x,m,label,col in [(6.02,'photo','Photo',MUTED),(6.61,'ours','Ours',TEAL)]:
        bh = ch*mean[m]/27
        d.rect(x,y0-bh,.30,bh,col,col,.1)
        d.text(x+.15,y0-bh-.10,f'{mean[m]:.2f}',8.5,col,bold=True,align='center')
        d.text(x+.15,1.81,label,7.7,align='center')
    d.text(6.34,2.00,f"+{mean['ours']-mean['photo']:.2f} dB",10,TEAL,bold=True,align='center')
    d.text(6.34,2.15,'8/8 scene PSNR wins',7.7,align='center')
    return d


def method(variant):
    refs = {
        'ral':('Sec. III-A · Eq. (1)','Sec. III-B · Eqs. (2–4)',
               'Sec. III-C · Eq. (5)','Sec. III-D · Eqs. (7–8)'),
        'master':('Ch. 2 · Eq. (2.7)','Ch. 4 · Eqs. (4.1–4.3)',
                  'Ch. 7 · Eq. (7.1)','Chs. 5–6 · Eq. (6.1)'),
    }[variant]
    d = Drawing(7.15,2.55)
    for x,label,ref in zip((.06,2.53,4.77),('(a) Predict','(b) Anchor','(c) Refine'),refs):
        d.heading(x,.13,label,ref)
    d.line([(.06,.46),(7.09,.46)],RULE,.5)
    d.image(ASSET/'target.png',.06,.69,.68)
    d.image(DATA/'desk/view1_gt.png',.06,1.36,.68)
    d.text(.39,.57,'Image pair',7.8,align='center')
    d.rect(1.06,1.00,.94,.69)
    d.text(1.53,1.23,'Shared\npredictor',8.3,bold=True,align='center')
    d.text(1.53,1.57,'Splatt3R',7.7,MUTED,align='center')
    d.line([(.74,.94),(.91,.94),(.91,1.17),(1.06,1.17)],arrow=True)
    d.line([(.74,1.61),(1.06,1.61)],arrow=True)
    d.rect(2.53,.67,1.65,.40)
    d.text(3.355,.87,'Tracker + pose graph',8.2,bold=True,align='center')
    d.line([(2.,1.17),(2.23,1.17),(2.23,.87),(2.53,.87)],BLUE,arrow=True)
    d.text(1.87,.72,'Geometry',7.5,BLUE,align='center')
    d.rect(2.53,1.44,1.65,.93,RULE,'white')
    d.text(2.61,1.28,r'Anchor $k$: $T_{W\mathcal{C}_k}$',8,
           ppt='Anchor k: T_WCₖ')
    d.line([(4.00,1.07),(4.00,1.44)],BLUE,arrow=True)
    gaussians(d,2.81,1.70,.64)
    d.text(3.16,1.70,r'Local map $\mathcal{G}_k$',8,ppt='Local map Gₖ')
    camera(d,2.64,2.14,.62)
    d.text(3.16,2.06,'RGB + relative',7.8)
    d.text(3.16,2.24,'camera pose',7.8)
    d.line([(2.,1.48),(2.20,1.48),(2.20,1.70),(2.53,1.70)],TEAL,arrow=True)
    d.text(1.63,1.85,'Gaussians',7.6,TEAL,align='center')
    d.line([(.74,1.76),(.84,1.76),(.84,2.17),(2.53,2.17)],arrow=True)
    d.text(1.52,2.07,'Tracked RGB',7.5,MUTED,align='center')
    d.rect(4.77,1.49,.89,.42)
    d.text(5.215,1.70,'Render',8.3,bold=True,align='center')
    d.line([(4.18,1.70),(4.77,1.70)],TEAL,arrow=True)
    d.rect(4.77,2.12,1.05,.40)
    d.text(5.295,2.32,'Sample view',8,align='center')
    d.line([(4.18,2.20),(4.49,2.20),(4.49,2.32),(4.77,2.32)],BLUE,arrow=True)
    d.line([(5.21,2.12),(5.21,1.91)],BLUE,arrow=True)
    d.rect(6.18,1.49,.89,.42)
    d.text(6.625,1.70,'Loss',8.3,bold=True,align='center')
    d.text(6.37,2.03,'L1 + SSIM',7.7,MUTED,align='center')
    d.line([(5.66,1.70),(6.18,1.70)],arrow=True)
    d.text(5.92,1.56,r'$\hat I$',8,ppt='Î',align='center')
    d.line([(5.82,2.32),(7.00,2.32),(7.00,1.91)],arrow=True)
    d.text(6.02,2.18,r'$I$',8,ppt='I',align='center')
    d.rect(6.18,.67,.89,.40)
    d.text(6.625,.87,'Adam',8.3,bold=True,align='center')
    d.line([(6.625,1.49),(6.625,1.07)],arrow=True)
    d.text(6.47,1.28,r'$\nabla_\theta\mathcal{L}$',8,ppt='∇θ L',align='right')
    d.line([(6.18,.87),(4.43,.87),(4.43,1.51),(4.18,1.51)],TEAL,arrow=True)
    d.text(5.22,.73,r'Update $\theta$',8,TEAL,ppt='Update θ',align='center')
    return d


def render_all(records):
    teaser(records).render('teaser')
    for variant in ('ral','master'):
        method(variant).render(f'method_overview_{variant}')
    from build_overview import build_overview
    build_overview()
    write_manifest(records)


def write_manifest(records):
    import hashlib
    drawings = {'teaser': teaser(records), 'method_overview_ral': method('ral'),
                'method_overview_master': method('master')}
    images = sorted({o['path'] for d in drawings.values()
                     for o in d.items if o['kind']=='image'})
    manifest = {
        'generator': 'scripts/paper/draw_paper_diagrams.py',
        'teaser': {'scene': 'office0', 'view': 'view0',
                   'selection': records['office0']['selection'],
                   'identical_crop_xywh': TEASER_ROI},
        'drawings': {name: {'size_inches': [d.width, d.height],
                            'minimum_label_pt': min(o['size'] for o in d.items
                                                     if o['kind']=='text'),
                            'continuous_arrows': sum(o.get('arrow',False)
                                                     for o in d.items)}
                     for name,d in drawings.items()},
        'figure_3': {
            'source': 'docs/Thesis/ral/fig/overview.tikz',
            'generator': 'scripts/paper/build_overview.py',
            'status': 'restored original system implementation figure',
        },
        'images': {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                   for p in images},
        'image_edits': 'uniform display scaling and identical explicitly marked crops only',
    }
    (FIG/'figures_1_3_manifest.json').write_text(json.dumps(manifest, indent=2)+'\n')


def editable_deck():
    from pptx import Presentation
    from pptx.util import Inches
    records = {s:json.loads((DATA/s/'metrics.json').read_text()) for s in SCENES}
    prs = Presentation()
    prs.slide_width,prs.slide_height = Inches(13.333),Inches(7.5)
    prs.core_properties.title = 'Splatt3R-SLAM — Figures 1–3'
    for drawing in (teaser(records), method('ral')):
        scale = 12.8/drawing.width
        drawing.add_slide(prs,x=.2665,y=(7.5-drawing.height*scale)/2,width=12.8)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(FIG/'overview.png'), Inches(.2665), Inches(.755),
        width=Inches(12.8), height=Inches(5.99))
    drawing = method('master')
    scale = 12.8/drawing.width
    drawing.add_slide(prs,x=.2665,y=(7.5-drawing.height*scale)/2,width=12.8)
    out = ROOT/'docs/Thesis/Splatt3R-SLAM-figures-1-3.pptx'
    prs.save(out)
    print(out)


if __name__=='__main__':
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pptx',action='store_true')
    args = parser.parse_args()
    if args.pptx:
        editable_deck()
    else:
        records = {s:json.loads((DATA/s/'metrics.json').read_text()) for s in SCENES}
        render_all(records)
